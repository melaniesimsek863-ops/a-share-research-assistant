from __future__ import annotations

from pathlib import Path
from typing import Mapping

import pandas as pd

from a_share_ai.models import (
    DATA_MODE_REAL_PUBLIC,
    DATA_MODE_SYNTHETIC_DEMO,
    DataQualityStatus,
)
from openpyxl.styles import Font, PatternFill


def write_excel_signals(
    candidates: pd.DataFrame,
    path: str | Path,
    *,
    metadata: Mapping[str, object] | None = None,
) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    data = candidates.copy()
    if "signal_labels" in data.columns:
        data["signal_labels"] = data["signal_labels"].map(
            lambda value: "、".join(value) if isinstance(value, list) else str(value)
        )
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        data_quality = metadata.get("data_quality") if metadata is not None else None
        if isinstance(data_quality, DataQualityStatus):
            banner = (
                "REAL PUBLIC MARKET DATA / \u516c\u5f00\u5e02\u573a\u6570\u636e"
                if data_quality.data_mode == DATA_MODE_REAL_PUBLIC
                else "SYNTHETIC DEMO DATA / \u5408\u6210\u6f14\u793a\u6570\u636e"
                if data_quality.data_mode == DATA_MODE_SYNTHETIC_DEMO
                else "DATA QUALITY METADATA / \u6570\u636e\u8d28\u91cf\u8bf4\u660e"
            )
            metadata_rows = [
                [banner, ""],
                ["provider", data_quality.provider],
                ["data_mode", data_quality.data_mode],
                ["latest_data_date", str(data_quality.latest_data_date or "")],
                ["report_date", str(data_quality.report_date)],
                ["adjustment", data_quality.adjustment],
                ["quality_status", data_quality.quality_status],
                ["warnings", ", ".join(data_quality.warnings)],
                ["blocking_reasons", ", ".join(data_quality.blocking_reasons)],
                ["data_warning", metadata.get("data_warning", "")],
                [
                    "uses_default_risk_state",
                    bool(metadata.get("uses_default_risk_state")),
                ],
                [
                    "\u98ce\u9669\u72b6\u6001",
                    "\u9ed8\u8ba4\u6f14\u793a\u98ce\u9669\u72b6\u6001\uff1b\u672a\u8fde\u63a5\u771f\u5b9e\u7ec4\u5408"
                    if metadata.get("uses_default_risk_state")
                    else "\u8c03\u7528\u65b9\u63d0\u4f9b\u7684\u98ce\u9669\u72b6\u6001",
                ],
                ["risk_state_source", metadata.get("risk_state_source", "unspecified")],
                ["risk_state", metadata.get("risk_state_summary", "unspecified")],
                ["\u6301\u4ed3\u7ba1\u7406\u8fb9\u754c", "\u6b62\u635f/\u79fb\u52a8\u6b62\u76c8\u53c2\u6570\u4ec5\u9884\u7559\uff0c\u5f53\u524d\u5019\u9009\u6d41\u7a0b\u672a\u6267\u884c\u9000\u51fa\u89c4\u5219"],
                ["exit_rule_status", "RESERVED / NOT ACTIVE"],
            ]
            pd.DataFrame(metadata_rows).to_excel(
                writer, index=False, header=False, sheet_name="\u8bf4\u660e"
            )
            metadata_sheet = writer.book["\u8bf4\u660e"]
            metadata_sheet["A1"].font = Font(bold=True, color="FFFFFF")
            metadata_sheet["A1"].fill = PatternFill("solid", fgColor="C00000")
            metadata_sheet.column_dimensions["A"].width = 22
            metadata_sheet.column_dimensions["B"].width = 72
        elif metadata is not None:
            is_demo = metadata.get("data_mode") == "synthetic_demo"
            banner = (
                "SYNTHETIC DEMO DATA / 合成演示数据"
                if is_demo
                else "DATA METADATA / 数据说明"
            )
            risk_state_text = (
                "默认演示风险状态；未连接真实组合"
                if metadata.get("uses_default_risk_state")
                else "调用方提供的风险状态"
            )
            metadata_rows = [
                [banner, ""],
                ["数据模式", metadata.get("data_mode", "unspecified")],
                ["最新数据日期", str(metadata.get("latest_data_date", "未知"))],
                ["数据警告", metadata.get("data_warning", "")],
                ["风险状态", risk_state_text],
                ["持仓管理边界", "止损/移动止盈参数仅预留，当前候选流程未执行退出规则"],
                ["risk_state_source", metadata.get("risk_state_source", "unspecified")],
                ["risk_state", metadata.get("risk_state_summary", "unspecified")],
                ["exit_rule_status", "RESERVED / NOT ACTIVE"],
            ]
            pd.DataFrame(metadata_rows).to_excel(
                writer, index=False, header=False, sheet_name="说明"
            )
            metadata_sheet = writer.book["说明"]
            metadata_sheet["A1"].font = Font(bold=True, color="FFFFFF")
            metadata_sheet["A1"].fill = PatternFill("solid", fgColor="C00000")
            metadata_sheet.column_dimensions["A"].width = 22
            metadata_sheet.column_dimensions["B"].width = 72
        data.to_excel(writer, index=False, sheet_name="signals")
    return output
