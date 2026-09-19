from datetime import date

import pandas as pd
from openpyxl import load_workbook

from a_share_ai.models import DataQualityStatus
from a_share_ai.reports.excel_report import write_excel_signals
from a_share_ai.reports.markdown_report import build_markdown_report


def test_build_markdown_report_contains_required_sections():
    candidates = pd.DataFrame(
        [
            {
                "code": "000001",
                "name": "强趋势",
                "tier": "focused",
                "total_score": 88.0,
                "suggested_position": 0.15,
                "risk_notes": "风险状态正常",
                "signal_labels": ["放量突破"],
            }
        ]
    )
    excluded = pd.DataFrame([{"code": "600001", "name": "ST测试", "exclude_reason": "st_or_delisting_risk"}])

    report = build_markdown_report(candidates, excluded, date(2026, 7, 29))

    assert "# A股短中线交易辅助日报 - 2026-07-29" in report
    assert "## 重点观察" in report
    assert "000001 强趋势" in report
    assert "所有交易必须人工确认" in report
    assert "## 剔除样例" in report


def test_write_excel_signals_creates_workbook(tmp_path):
    candidates = pd.DataFrame(
        [
            {
                "code": "000001",
                "name": "强趋势",
                "tier": "focused",
                "total_score": 88.0,
                "suggested_position": 0.15,
                "can_buy": True,
                "risk_action": "normal",
                "risk_notes": "风险状态正常",
                "signal_labels": ["放量突破"],
            }
        ]
    )
    output = write_excel_signals(candidates, tmp_path / "signals.xlsx")

    workbook = load_workbook(output)
    sheet = workbook["signals"]
    assert sheet["A1"].value == "code"
    assert sheet["A2"].value == "000001"
    assert sheet["I2"].value == "放量突破"


def test_markdown_report_discloses_demo_data_and_inactive_exit_controls():
    candidates = pd.DataFrame(
        [
            {
                "code": "000001",
                "name": "演示标的",
                "tier": "focused",
                "total_score": 88.0,
                "suggested_position": 0.15,
                "risk_notes": "风险状态正常",
                "signal_labels": ["20日涨幅低于18%"],
            }
        ]
    )
    excluded = pd.DataFrame()

    report = build_markdown_report(
        candidates,
        excluded,
        date(2026, 7, 29),
        data_mode="synthetic_demo",
        latest_data_date=date(2026, 7, 29),
        data_warning="仅用于离线流程演示。",
        uses_default_risk_state=True,
    )

    assert "SYNTHETIC DEMO DATA / 合成演示数据" in report
    assert "最新数据日期：2026-07-29" in report
    assert "默认演示风险状态" in report
    assert "止损/移动止盈参数仅为未来持仓管理模块预留" in report
    assert "EXIT RULE STATUS: RESERVED / NOT ACTIVE" in report


def test_write_excel_signals_adds_conspicuous_demo_metadata_sheet(tmp_path):
    candidates = pd.DataFrame(
        [
            {
                "code": "000001",
                "name": "演示标的",
                "tier": "focused",
                "can_buy": False,
            }
        ]
    )

    output = write_excel_signals(
        candidates,
        tmp_path / "signals.xlsx",
        metadata={
            "data_mode": "synthetic_demo",
            "latest_data_date": "2026-07-29",
            "data_warning": "报告日期晚于演示数据日期，不生成可买建议。",
            "uses_default_risk_state": True,
        },
    )

    workbook = load_workbook(output)
    sheet = workbook["说明"]
    assert sheet["A1"].value == "SYNTHETIC DEMO DATA / 合成演示数据"
    assert sheet["A3"].value == "最新数据日期"
    assert sheet["B3"].value == "2026-07-29"
    assert "不生成可买建议" in sheet["B4"].value


def test_markdown_report_shows_real_public_data_banner() -> None:
    status = DataQualityStatus(
        provider="akshare",
        data_mode="real_public_market_data",
        latest_data_date=date(2026, 7, 28),
        report_date=date(2026, 7, 29),
        is_stale=False,
        is_future_dated=False,
        quality_status="ok",
        warnings=[],
        blocking_reasons=[],
        adjustment="qfq",
    )

    text = build_markdown_report(
        pd.DataFrame(columns=["tier"]),
        pd.DataFrame(),
        date(2026, 7, 29),
        data_quality=status,
    )

    assert "REAL PUBLIC MARKET DATA / \u516c\u5f00\u5e02\u573a\u6570\u636e" in text
    assert "provider: akshare" in text
    assert "\u590d\u6743\u65b9\u5f0f\uff1aqfq" in text


def test_excel_metadata_sheet_contains_quality_fields(tmp_path) -> None:
    status = DataQualityStatus(
        provider="akshare",
        data_mode="real_public_market_data",
        latest_data_date=date(2026, 7, 28),
        report_date=date(2026, 7, 29),
        is_stale=False,
        is_future_dated=False,
        quality_status="ok",
        warnings=[],
        blocking_reasons=[],
        adjustment="qfq",
    )
    path = tmp_path / "signals.xlsx"

    write_excel_signals(pd.DataFrame(), path, metadata={"data_quality": status})

    workbook = load_workbook(path)
    rows = list(workbook["\u8bf4\u660e"].iter_rows(values_only=True))
    labels = {row[0]: row[1] for row in rows if row[0]}

    assert rows[0][0] == "REAL PUBLIC MARKET DATA / \u516c\u5f00\u5e02\u573a\u6570\u636e"
    assert labels["provider"] == "akshare"
    assert labels["adjustment"] == "qfq"
    assert labels["quality_status"] == "ok"
