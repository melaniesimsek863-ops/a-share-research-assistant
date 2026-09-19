from datetime import date
import inspect

import pandas as pd
import pytest
from openpyxl import load_workbook

from tests.fixtures.sample_data import sample_price_history, sample_stock_info

import a_share_ai.pipeline as pipeline
import a_share_ai.data.providers as providers
import scripts.smoke_akshare_limited as limited_smoke
from a_share_ai.data.cache import CsvMarketCache
from a_share_ai.data.cached_provider import CachedMarketDataProvider
from a_share_ai.data.providers import DataProviderError, LocalSampleProvider
from a_share_ai.models import (
    DATA_MODE_REAL_PUBLIC,
    QUALITY_OK,
    DataQualityStatus,
    RiskState,
)
from a_share_ai.pipeline import run_daily_pipeline

class BrokenRealProvider:
    data_mode = "real_public_market_data"
    provider_name = "akshare"

    def get_stock_info(self):
        raise DataProviderError("akshare stock_info failed: remote changed")


class EmptyHistoryProvider:
    data_mode = DATA_MODE_REAL_PUBLIC
    provider_name = "empty-history-fake"

    def get_stock_info(self) -> pd.DataFrame:
        return pd.DataFrame(
            [
                {
                    "code": "000001",
                    "name": "\u5e73\u5b89\u94f6\u884c",
                    "board": "main",
                    "is_st": False,
                    "is_delisting_risk": False,
                    "listing_date": date(1991, 4, 3),
                    "is_suspended": False,
                    "avg_turnover_20d": float("nan"),
                }
            ]
        )

    def get_price_history(self, codes: list[str]) -> pd.DataFrame:
        assert codes == ["000001"]
        return pd.DataFrame(columns=["date", "code", "open", "high", "low", "close", "volume", "amount"])

    def get_index_history(self) -> pd.DataFrame:
        return pd.DataFrame(columns=["date", "code", "open", "high", "low", "close", "volume", "amount"])


def test_pipeline_does_not_create_buyable_candidates_when_history_is_empty(tmp_path) -> None:
    result = run_daily_pipeline(
        "configs/default.yaml",
        date(2026, 7, 30),
        provider=EmptyHistoryProvider(),
        output_dir=tmp_path,
        risk_state=RiskState(current_drawdown=-0.03, current_holdings=0, daily_new_buys=0),
    )

    candidates = result["candidates"]
    assert candidates.empty or not candidates["can_buy"].fillna(False).any()
    assert result["data_quality"].quality_status == "blocked"
    assert "no_price_data" in result["data_quality"].blocking_reasons


def test_pipeline_does_not_buy_when_cached_provider_has_no_remote_or_cache(tmp_path) -> None:
    class EmptyCachedRealProvider:
        data_mode = DATA_MODE_REAL_PUBLIC
        provider_name = "akshare-cached"

        def get_stock_info(self) -> pd.DataFrame:
            return pd.DataFrame(
                [{
                    "code": "000001",
                    "name": "bank",
                    "board": "main",
                    "listing_date": date(1991, 4, 3),
                    "is_st": False,
                    "is_delisting_risk": False,
                    "is_suspended": False,
                    "avg_turnover_20d": float("nan"),
                }]
            )

        def get_price_history(self, codes: list[str]) -> pd.DataFrame:
            return pd.DataFrame(columns=["date", "code", "open", "high", "low", "close", "volume", "amount"])

        def get_index_history(self) -> pd.DataFrame:
            return pd.DataFrame(columns=["date", "code", "open", "high", "low", "close", "volume", "amount"])

    result = run_daily_pipeline(
        "configs/default.yaml",
        date(2026, 7, 30),
        provider=EmptyCachedRealProvider(),
        output_dir=tmp_path,
        risk_state=RiskState(current_drawdown=-0.03, current_holdings=0, daily_new_buys=0),
    )

    assert result["data_quality"].quality_status == "blocked"
    assert result["candidates"].empty or not result["candidates"]["can_buy"].fillna(False).any()


def test_pipeline_blocks_candidates_when_one_cached_symbol_refresh_fails(tmp_path) -> None:
    report_date = date(2026, 7, 30)

    class MixedRefreshProvider:
        data_mode = DATA_MODE_REAL_PUBLIC
        provider_name = "mixed-refresh"

        def get_stock_info(self) -> pd.DataFrame:
            return sample_stock_info().query("code in ['000001', '300001']").copy()

        def get_price_history(self, codes: list[str]) -> pd.DataFrame:
            assert len(codes) == 1
            if codes[0] == "000001":
                raise DataProviderError("remote timeout")
            healthy = sample_price_history("300001")
            healthy["date"] = pd.to_datetime(healthy["date"]) + pd.Timedelta(days=1)
            healthy["date"] = healthy["date"].dt.date
            return healthy

        def get_index_history(self) -> pd.DataFrame:
            return pd.DataFrame()

    cache = CsvMarketCache(tmp_path / "akshare")
    cache.write_daily("000001", sample_price_history("000001"))
    provider = CachedMarketDataProvider(
        MixedRefreshProvider(), cache, report_date=report_date
    )

    result = run_daily_pipeline(
        "configs/default.yaml",
        report_date,
        provider=provider,
        output_dir=tmp_path / "outputs",
        risk_state=RiskState(current_drawdown=-0.03, current_holdings=0, daily_new_buys=0),
    )

    assert "000001" in result["candidates"]["code"].tolist()
    assert result["data_quality"].quality_status == "blocked"
    assert "cache_refresh_failed" in result["data_quality"].blocking_reasons
    assert not result["candidates"]["can_buy"].fillna(False).any()

def test_limited_smoke_rewrites_outputs_as_observation_only(tmp_path) -> None:
    """Removing the smoke guard must never leave a buyable signal in either output."""
    assert hasattr(limited_smoke, "finalize_limited_smoke")
    report_date = date(2026, 7, 30)
    result = {
        "candidates": pd.DataFrame([{
            "code": "000001", "name": "bank", "tier": "focused", "total_score": 1.0,
            "signal_labels": [], "suggested_position": 0.15, "risk_notes": "",
            "can_buy": True, "risk_action": "buy_allowed",
        }]),
        "excluded": pd.DataFrame(columns=["code", "name", "exclude_reason"]),
        "markdown_path": tmp_path / "reports" / "smoke.md",
        "excel_path": tmp_path / "signals" / "smoke.xlsx",
        "data_quality": DataQualityStatus(
            provider="akshare-limited-smoke",
            data_mode=DATA_MODE_REAL_PUBLIC,
            latest_data_date=report_date,
            report_date=report_date,
            is_stale=False,
            is_future_dated=False,
            quality_status=QUALITY_OK,
        ),
        "data_warning": None,
        "risk_state": RiskState(
            current_drawdown=-0.03, current_holdings=0, daily_new_buys=0
        ),
    }

    updated = limited_smoke.finalize_limited_smoke(result)

    assert not updated["candidates"]["can_buy"].fillna(False).any()
    assert (updated["candidates"]["suggested_position"] == 0).all()
    assert (
        updated["candidates"]["risk_action"] == "limited_smoke_observe_only"
    ).all()
    assert "LIMITED OPERATIONAL SMOKE / NOT A RECOMMENDATION" in (
        updated["markdown_path"].read_text(encoding="utf-8")
    )
    signals = pd.read_excel(updated["excel_path"], sheet_name="signals")
    assert not signals["can_buy"].fillna(False).any()
    assert (signals["suggested_position"] == 0).all()
    assert (signals["risk_action"] == "limited_smoke_observe_only").all()
    workbook = load_workbook(updated["excel_path"])
    assert workbook[workbook.sheetnames[0]]["A16"].value == "LIMITED OPERATIONAL SMOKE / NOT A RECOMMENDATION"

def test_pipeline_forces_limited_smoke_observation_only_before_first_output_write(
    monkeypatch, tmp_path
) -> None:
    """An interrupted limited smoke must not leave a buyable first-write report."""
    def interrupt_excel_write(*args, **kwargs) -> None:
        raise RuntimeError("simulated Excel write interruption")

    monkeypatch.setattr(pipeline, "write_excel_signals", interrupt_excel_write)

    with pytest.raises(RuntimeError, match="simulated Excel write interruption"):
        run_daily_pipeline(
            "configs/default.yaml",
            date(2026, 7, 29),
            provider=LocalSampleProvider(),
            output_dir=tmp_path,
            force_observation_only=True,
        )

    report = (tmp_path / "reports" / "2026-07-29_daily_report.md").read_text(
        encoding="utf-8"
    )
    assert "\u53ef\u4e70\u5efa\u8bae\uff1a\u662f\uff08\u4ecd\u9700\u4eba\u5de5\u786e\u8ba4\uff09" not in report
    assert "\u53ef\u4e70\u5efa\u8bae\uff1a\u5426\uff08\u4ec5\u89c2\u5bdf\uff09" in report


def test_pipeline_real_provider_failure_is_not_silent(tmp_path):
    try:
        run_daily_pipeline(
            "configs/default.yaml",
            date(2026, 7, 29),
            provider=BrokenRealProvider(),
            output_dir=tmp_path,
        )
    except DataProviderError as exc:
        assert "akshare stock_info failed" in str(exc)
    else:
        raise AssertionError("expected DataProviderError")


class FutureAndStaleProvider:
    data_mode = "real_public_market_data"
    provider_name = "akshare"

    def get_stock_info(self):
        return pd.DataFrame(
            [
                {
                    "code": "000001",
                    "name": "平安银行",
                    "board": "main",
                    "listing_date": date(1991, 4, 3),
                    "is_st": False,
                    "is_delisting_risk": False,
                    "is_suspended": False,
                    "avg_turnover_20d": 500000000,
                }
            ]
        )

    def get_price_history(self, codes):
        return pd.DataFrame(
            [
                {"date": date(2026, 7, 20), "code": "000001", "open": 10, "high": 11, "low": 9, "close": 10, "volume": 100, "amount": 1000},
                {"date": date(2026, 7, 30), "code": "000001", "open": 10, "high": 11, "low": 9, "close": 10, "volume": 100, "amount": 1000},
            ]
        )

    def get_index_history(self):
        return pd.DataFrame()


def test_pipeline_blocks_buy_when_real_data_is_stale(tmp_path):
    result = run_daily_pipeline(
        "configs/default.yaml",
        date(2026, 7, 29),
        provider=FutureAndStaleProvider(),
        output_dir=tmp_path,
    )

    assert result["data_quality"].quality_status == "blocked"
    assert result["data_quality"].is_future_dated is True
    assert result["candidates"]["can_buy"].eq(False).all()

def test_production_provider_module_does_not_import_tests_package():
    source = inspect.getsource(providers)
    assert "tests.fixtures" not in source


def test_run_daily_pipeline_outputs_report_and_excel(tmp_path):
    result = run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=date(2026, 7, 29),
        provider=LocalSampleProvider(),
        output_dir=tmp_path,
    )

    assert result["markdown_path"].exists()
    assert result["excel_path"].exists()
    assert "candidates" in result
    assert not result["candidates"].empty
    assert "excluded" in result
    assert not result["excluded"].empty
    assert date(2026, 7, 29).isoformat() in result["markdown_path"].read_text(encoding="utf-8")


def test_run_daily_pipeline_does_not_use_prices_after_report_date(tmp_path):
    report_date = date(2026, 7, 15)

    result = run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=report_date,
        provider=LocalSampleProvider(),
        output_dir=tmp_path,
    )

    assert result["candidates"]["date"].max() <= report_date
    assert report_date.isoformat() in result["markdown_path"].read_text(encoding="utf-8")


def test_run_daily_pipeline_outputs_empty_files_for_empty_eligible_universe(tmp_path):
    class EmptyUniverseProvider(LocalSampleProvider):
        def get_stock_info(self):
            stocks = sample_stock_info()
            stocks["is_st"] = True
            return stocks

    provider = EmptyUniverseProvider()
    result = run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=date(2026, 7, 29),
        provider=provider,
        output_dir=tmp_path,
    )

    assert provider.get_price_history([]).empty
    assert result["candidates"].empty
    assert not result["excluded"].empty
    assert result["markdown_path"].exists()
    assert result["excel_path"].exists()


def test_run_daily_pipeline_outputs_backtest_summary(tmp_path):
    result = run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=date(2026, 7, 29),
        provider=LocalSampleProvider(),
        output_dir=tmp_path,
    )

    assert result["backtest_path"].exists()
    text = result["backtest_path"].read_text(encoding="utf-8")
    assert "# 基础回测摘要" in text
    assert "最大回撤" in text


def test_run_daily_pipeline_applies_supplied_risk_state_sequentially(tmp_path):
    result = run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=date(2026, 7, 29),
        provider=LocalSampleProvider(),
        output_dir=tmp_path,
        risk_state=RiskState(
            current_drawdown=-0.05,
            current_holdings=5,
            daily_new_buys=2,
        ),
    )

    assert result["candidates"]["can_buy"].tolist() == [True, False]
    assert result["candidates"]["risk_action"].tolist() == ["normal", "hold_limit"]


def test_default_demo_pipeline_labels_outputs_with_latest_data_date(tmp_path):
    result = run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=date(2026, 7, 29),
        output_dir=tmp_path,
    )

    assert result["data_mode"] == "synthetic_demo"
    assert result["latest_data_date"] == date(2026, 7, 29)
    markdown = result["markdown_path"].read_text(encoding="utf-8")
    backtest = result["backtest_path"].read_text(encoding="utf-8")
    assert "SYNTHETIC DEMO DATA / 合成演示数据" in markdown
    assert "最新数据日期：2026-07-29" in markdown
    assert "SYNTHETIC DEMO DATA / 合成演示数据" in backtest



def test_run_daily_pipeline_uses_provider_override_with_config_adjustment(
    monkeypatch, tmp_path
):
    """A provider-name override must drive selection without contacting the network."""
    captured = {}

    def fake_create_provider(
        name,
        adjustment="qfq",
        *,
        cache_enabled=False,
        cache_base_dir="data/raw/akshare",
        report_date=None,
    ):
        captured["name"] = name
        captured["adjustment"] = adjustment
        captured["cache_enabled"] = cache_enabled
        captured["cache_base_dir"] = cache_base_dir
        captured["report_date"] = report_date
        return LocalSampleProvider()

    monkeypatch.setattr(pipeline, "create_provider", fake_create_provider)

    pipeline.run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=date(2026, 7, 29),
        provider_name="akshare",
        output_dir=tmp_path,
    )

    assert captured == {
        "name": "akshare",
        "adjustment": "qfq",
        "cache_enabled": False,
        "cache_base_dir": "data/raw/akshare",
        "report_date": date(2026, 7, 29),
    }
def test_future_dated_demo_report_emits_no_buyable_suggestions(tmp_path):
    result = run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=date(2026, 7, 30),
        provider=LocalSampleProvider(),
        output_dir=tmp_path,
    )

    assert not result["candidates"]["can_buy"].any()
    assert (result["candidates"]["suggested_position"] == 0.0).all()
    assert (result["candidates"]["risk_action"] == "stale_demo_data").all()
    assert result["latest_data_date"] == date(2026, 7, 29)
    report = result["markdown_path"].read_text(encoding="utf-8")
    assert "报告日期晚于合成演示数据最新日期" in report
    assert "不生成可买建议" in report
    assert "可买建议：否（仅观察）" in report


def test_pipeline_future_rows_are_blocking_even_when_also_filtered(tmp_path):
    result = run_daily_pipeline(
        "configs/default.yaml",
        date(2026, 7, 29),
        provider=FutureAndStaleProvider(),
        output_dir=tmp_path,
    )

    assert "future_data_detected" in result["data_quality"].blocking_reasons
    assert result["data_quality"].quality_status == "blocked"


def test_pipeline_blocks_invalid_future_row_removed_by_bar_validation(tmp_path):
    class InvalidFutureRowProvider(LocalSampleProvider):
        data_mode = "real_public_market_data"
        provider_name = "invalid-future-row"

        def get_price_history(self, codes):
            valid_history = super().get_price_history(codes)
            invalid_future = valid_history.iloc[[-1]].copy()
            invalid_future["date"] = date(2026, 7, 30)
            invalid_future["volume"] = "invalid"
            return pd.concat([valid_history, invalid_future], ignore_index=True)

    report_date = date(2026, 7, 29)
    result = run_daily_pipeline(
        "configs/default.yaml",
        report_date,
        provider=InvalidFutureRowProvider(),
        output_dir=tmp_path,
    )

    assert result["data_quality"].quality_status == "blocked"
    assert result["data_quality"].is_future_dated is True
    assert "future_data_detected" in result["data_quality"].blocking_reasons
    assert "future_data_removed" in result["data_quality"].warnings
    assert "invalid_bar_rows_removed:1" in result["data_quality"].warnings
    assert result["candidates"]["date"].max() <= report_date
    assert not result["candidates"]["can_buy"].any()


def test_pipeline_blocks_malformed_bars_before_features(tmp_path):
    class MalformedBarsProvider(LocalSampleProvider):
        data_mode = "real_public_market_data"
        provider_name = "malformed"

        def get_price_history(self, codes):
            return pd.DataFrame(
                [{"date": date(2026, 7, 29), "code": codes[0], "close": 10.0}]
            )

    result = run_daily_pipeline(
        "configs/default.yaml",
        date(2026, 7, 29),
        provider=MalformedBarsProvider(),
        output_dir=tmp_path,
    )

    assert result["data_quality"].quality_status == "blocked"
    assert "missing_bar_columns:amount,high,low,open,volume" in (
        result["data_quality"].blocking_reasons
    )
    assert result["candidates"].empty or not result["candidates"]["can_buy"].any()


def test_pipeline_derives_20d_turnover_before_final_universe_filter(tmp_path):
    class HistoryLiquidityProvider(LocalSampleProvider):
        data_mode = "real_public_market_data"
        provider_name = "history-liquidity"

        def get_stock_info(self):
            stocks = super().get_stock_info().iloc[:2].copy()
            stocks["avg_turnover_20d"] = float("nan")
            return stocks

        def get_price_history(self, codes):
            frames = []
            for code in codes:
                frame = super().get_price_history([code])
                frame["amount"] = 200_000_000 if code == "000001" else 20_000_000
                frames.append(frame)
            return pd.concat(frames, ignore_index=True)

    result = run_daily_pipeline(
        "configs/default.yaml",
        date(2026, 7, 29),
        provider=HistoryLiquidityProvider(),
        output_dir=tmp_path,
    )

    assert result["candidates"]["code"].tolist() == ["000001"]
    assert "300001" in result["excluded"]["code"].tolist()


def test_pipeline_excel_keeps_risk_safety_metadata_with_data_quality(tmp_path):
    result = run_daily_pipeline(
        config_path="configs/default.yaml",
        report_date=date(2026, 7, 29),
        provider=LocalSampleProvider(),
        output_dir=tmp_path,
    )

    workbook = load_workbook(result["excel_path"])
    rows = list(workbook["\u8bf4\u660e"].iter_rows(values_only=True))
    labels = {row[0]: row[1] for row in rows if row[0]}

    assert labels["quality_status"] == "ok"
    assert labels["risk_state_source"] == "default_demo"
    assert labels["risk_state"] == (
        "current_drawdown=0.0, current_holdings=0, daily_new_buys=0"
    )
    assert labels["uses_default_risk_state"] is True
    assert labels["\u98ce\u9669\u72b6\u6001"] == (
        "\u9ed8\u8ba4\u6f14\u793a\u98ce\u9669\u72b6\u6001\uff1b\u672a\u8fde\u63a5\u771f\u5b9e\u7ec4\u5408"
    )
    assert labels["exit_rule_status"] == "RESERVED / NOT ACTIVE"
    assert "\u5f53\u524d\u5019\u9009\u6d41\u7a0b\u672a\u6267\u884c\u9000\u51fa\u89c4\u5219" in (
        labels["\u6301\u4ed3\u7ba1\u7406\u8fb9\u754c"]
    )


def test_pipeline_blocks_bars_with_unusable_numeric_values(tmp_path):
    class NonNumericBarsProvider(LocalSampleProvider):
        data_mode = "real_public_market_data"
        provider_name = "non-numeric"

        def get_price_history(self, codes):
            return pd.DataFrame(
                [{
                    "date": date(2026, 7, 29),
                    "code": codes[0],
                    "open": 10,
                    "high": 11,
                    "low": 9,
                    "close": 10,
                    "volume": "not-a-number",
                    "amount": None,
                }]
            )

    result = run_daily_pipeline(
        "configs/default.yaml",
        date(2026, 7, 29),
        provider=NonNumericBarsProvider(),
        output_dir=tmp_path,
    )

    assert result["data_quality"].quality_status == "blocked"
    assert "unusable_bar_columns:amount,volume" in (
        result["data_quality"].blocking_reasons
    )
    assert result["candidates"].empty or not result["candidates"]["can_buy"].any()
