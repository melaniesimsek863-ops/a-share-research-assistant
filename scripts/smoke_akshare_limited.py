from __future__ import annotations

import argparse
import sys
from dataclasses import replace
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from a_share_ai.data.akshare_provider import AkShareProvider
from a_share_ai.data.cache import CsvMarketCache
from a_share_ai.data.cached_provider import CachedMarketDataProvider
from a_share_ai.data.providers import DataProviderError
from a_share_ai.diagnostics.coverage import (
    CoverageSummary,
    render_coverage_terminal_summary,
    summarize_manifest,
    write_coverage_report,
)
from a_share_ai.models import DATA_MODE_REAL_PUBLIC, RiskState
from a_share_ai.pipeline import run_daily_pipeline
from a_share_ai.reports.excel_report import write_excel_signals
from a_share_ai.reports.markdown_report import build_markdown_report

LIMITED_SMOKE_NOTICE = "LIMITED OPERATIONAL SMOKE / NOT A RECOMMENDATION"


def default_coverage_output_dir(limit: int) -> Path:
    return Path("outputs") / "smoke-akshare-limited" / f"coverage-limit{limit}"


def resolve_coverage_output_dir(
    limit: int, explicit_output_dir: str | Path | None
) -> Path:
    if explicit_output_dir is not None:
        return Path(explicit_output_dir)
    return default_coverage_output_dir(limit)


class LimitedProvider:
    data_mode = DATA_MODE_REAL_PUBLIC
    provider_name = "akshare-limited-smoke"

    def __init__(self, inner: AkShareProvider, limit: int) -> None:
        self._inner = inner
        self._limit = limit

    def get_stock_info(self) -> pd.DataFrame:
        stock_info = self._inner.get_stock_info()
        return stock_info.head(self._limit).reset_index(drop=True)

    def get_price_history(self, codes: list[str]) -> pd.DataFrame:
        return self._inner.get_price_history(codes)

    def get_index_history(self) -> pd.DataFrame:
        return self._inner.get_index_history()


def build_limited_provider(
    limit: int,
    report_date: date,
    cache_enabled: bool = False,
    cache_dir: str | Path = "data/raw/akshare-smoke",
) -> object:
    limited = LimitedProvider(AkShareProvider(), limit)
    if not cache_enabled:
        return limited
    return CachedMarketDataProvider(
        limited,
        CsvMarketCache(cache_dir),
        report_date=report_date,
    )


def finalize_limited_smoke(result: dict[str, Any]) -> dict[str, Any]:
    """Rewrite limited smoke outputs so a truncated universe never looks buyable."""
    candidates = result["candidates"].copy()
    if not candidates.empty:
        candidates["can_buy"] = False
        candidates["suggested_position"] = 0.0
        candidates["risk_action"] = "limited_smoke_observe_only"
        existing_notes = candidates.get(
            "risk_notes", pd.Series([""] * len(candidates), index=candidates.index)
        ).fillna("")
        candidates["risk_notes"] = existing_notes.map(
            lambda value: f"{value}; {LIMITED_SMOKE_NOTICE}".strip("; ")
        )
    result["candidates"] = candidates

    data_quality = result["data_quality"]
    report_date = data_quality.report_date
    markdown = build_markdown_report(
        candidates,
        result["excluded"],
        report_date,
        data_mode=data_quality.data_mode,
        latest_data_date=data_quality.latest_data_date,
        data_warning=LIMITED_SMOKE_NOTICE,
        data_quality=data_quality,
        uses_default_risk_state=False,
        risk_state=result.get("risk_state"),
    )
    markdown_path = Path(result["markdown_path"])
    markdown_path.parent.mkdir(parents=True, exist_ok=True)
    markdown_path.write_text(markdown, encoding="utf-8")

    excel_path = Path(result["excel_path"])
    write_excel_signals(
        candidates,
        excel_path,
        metadata={
            "data_quality": data_quality,
            "data_mode": data_quality.data_mode,
            "latest_data_date": data_quality.latest_data_date,
            "data_warning": LIMITED_SMOKE_NOTICE,
            "uses_default_risk_state": False,
            "risk_state_source": "limited_smoke",
            "risk_state_summary": str(result.get("risk_state", "unspecified")),
        },
    )
    workbook = load_workbook(excel_path)
    notice_sheet = workbook[workbook.sheetnames[0]]
    notice_sheet["A16"] = LIMITED_SMOKE_NOTICE
    notice_sheet["A16"].font = Font(bold=True, color="FFFFFF")
    notice_sheet["A16"].fill = PatternFill("solid", fgColor="C00000")
    workbook.save(excel_path)
    return result


def build_pipeline_funnel_counts(result: dict[str, Any]) -> dict[str, int]:
    candidates = result.get("candidates", pd.DataFrame())
    excluded = result.get("excluded", pd.DataFrame())
    if not isinstance(candidates, pd.DataFrame):
        candidates = pd.DataFrame()
    if not isinstance(excluded, pd.DataFrame):
        excluded = pd.DataFrame()

    if "exclude_reason" in excluded.columns:
        reasons = excluded["exclude_reason"].fillna("").astype(str)
        pre_fetch_excluded = int(reasons.eq("st_or_delisting_risk").sum())
        post_fetch_excluded = int(len(excluded) - pre_fetch_excluded)
    else:
        pre_fetch_excluded = 0
        post_fetch_excluded = len(excluded)

    return {
        "pre_fetch_excluded": pre_fetch_excluded,
        "post_fetch_excluded": post_fetch_excluded,
        "final_candidates": len(candidates),
    }


def write_partial_coverage_report(
    *,
    report_date: date,
    requested_limit: int,
    cache_dir: str | Path,
    output_dir: str | Path,
    partial_status: str,
    partial_error: str = "",
) -> tuple[Path, CoverageSummary]:
    coverage_path, summary = write_coverage_report_from_manifest(
        report_date=report_date,
        requested_limit=requested_limit,
        cache_dir=cache_dir,
        output_dir=output_dir,
        smoke_status=partial_status,
        smoke_error=partial_error,
    )
    blocked_summary = replace(
        summary,
        blocked=True,
        blocking_reason=partial_status,
    )
    coverage_path = write_coverage_report(blocked_summary, output_dir)
    return coverage_path, blocked_summary


def write_coverage_report_from_manifest(
    *,
    report_date: date,
    requested_limit: int,
    cache_dir: str | Path,
    output_dir: str | Path = Path("outputs") / "smoke-akshare-limited" / "coverage",
    smoke_status: str = "unknown",
    smoke_error: str = "",
    funnel_counts: dict[str, int] | None = None,
) -> tuple[Path, CoverageSummary]:
    cache = CsvMarketCache(cache_dir)
    manifest_path = cache.manifest_path
    manifest = cache.read_manifest() if manifest_path.exists() else None
    summary = summarize_manifest(
        manifest,
        report_date=report_date,
        requested_limit=requested_limit,
        cache_dir=cache_dir,
        manifest_path=manifest_path,
        smoke_status=smoke_status,
        smoke_error=smoke_error,
        funnel_counts=funnel_counts,
    )
    return write_coverage_report(summary, output_dir), summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Limited AkShare real-data smoke test")
    parser.add_argument("--config", default="configs/default.yaml")
    parser.add_argument("--date", required=True)
    parser.add_argument("--limit", type=int, default=20)
    parser.add_argument("--cache-enabled", action="store_true")
    parser.add_argument("--cache-dir", default="data/raw/akshare-smoke")
    parser.add_argument("--coverage-report", action="store_true")
    parser.add_argument("--coverage-output-dir")
    parser.add_argument("--partial-coverage-only", action="store_true")
    parser.add_argument("--partial-status", default="timeout_partial")
    parser.add_argument("--partial-error", default="")
    args = parser.parse_args()

    if args.limit <= 0:
        parser.exit(status=2, message="--limit must be positive\n")

    report_date = date.fromisoformat(args.date)
    coverage_output_dir = resolve_coverage_output_dir(
        args.limit, args.coverage_output_dir
    )
    if args.partial_coverage_only:
        coverage_path, coverage_summary = write_partial_coverage_report(
            report_date=report_date,
            requested_limit=args.limit,
            cache_dir=args.cache_dir,
            output_dir=coverage_output_dir,
            partial_status=args.partial_status,
            partial_error=args.partial_error,
        )
        print(render_coverage_terminal_summary(coverage_summary))
        print(f"Coverage report: {coverage_path}")
        return

    provider = build_limited_provider(
        args.limit,
        report_date,
        cache_enabled=args.cache_enabled,
        cache_dir=args.cache_dir,
    )
    try:
        result = run_daily_pipeline(
            args.config,
            report_date,
            provider=provider,
            output_dir=Path("outputs") / "smoke-akshare-limited",
            risk_state=RiskState(
                current_drawdown=-0.03, current_holdings=0, daily_new_buys=0
            ),
            force_observation_only=True,
        )
    except DataProviderError as exc:
        if args.coverage_report:
            try:
                coverage_path, coverage_summary = write_coverage_report_from_manifest(
                    report_date=report_date,
                    requested_limit=args.limit,
                    cache_dir=args.cache_dir,
                    smoke_status="provider_failed",
                    smoke_error=str(exc),
                    output_dir=coverage_output_dir,
                )
            except Exception as coverage_exc:  # noqa: BLE001 - preserve provider error even if diagnostics fail
                print(f"Coverage report failed: {coverage_exc}", file=sys.stderr)
            else:
                print(render_coverage_terminal_summary(coverage_summary))
                print(f"Coverage report: {coverage_path}")
        parser.exit(status=2, message=f"Data provider failed: {exc}\n")
    result = finalize_limited_smoke(result)
    print(f"Markdown report: {result['markdown_path']}")
    print(f"Excel signals: {result['excel_path']}")
    print(f"Data quality: {result['data_quality'].quality_status}")
    print(LIMITED_SMOKE_NOTICE)
    if args.cache_enabled:
        manifest_path = Path(args.cache_dir) / "manifest.csv"
        print(f"Cache enabled: {args.cache_dir}")
        print(f"Manifest: {manifest_path}")
        if manifest_path.exists():
            manifest = pd.read_csv(manifest_path)
            status_counts = (
                manifest.get("status", pd.Series(dtype=object)).value_counts().to_dict()
            )
            print(f"Manifest status counts: {status_counts}")
    if args.coverage_report:
        coverage_path, coverage_summary = write_coverage_report_from_manifest(
            report_date=report_date,
            requested_limit=args.limit,
            cache_dir=args.cache_dir,
            smoke_status="completed",
            output_dir=coverage_output_dir,
            funnel_counts=build_pipeline_funnel_counts(result),
        )
        print(render_coverage_terminal_summary(coverage_summary))
        print(f"Coverage report: {coverage_path}")


if __name__ == "__main__":
    main()
